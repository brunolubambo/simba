"""Documentos: transforma texto em markdown simples num arquivo .docx (editável) ou .pdf no workspace,
pronto para mandar pelo Telegram ou anexar. Pensado para CV, carta de apresentação, planos e relatórios."""
import re, unicodedata
from pathlib import Path
from claude_agent_sdk import tool, create_sdk_mcp_server
from .config import WORKSPACE
from .toolkit import schema, safe

OUT = WORKSPACE / "documentos"


def _name(nome: str) -> str:
    s = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode()
    s = re.sub(r"[^\w\- ]+", "", s).strip().replace(" ", "_")
    return s[:60] or "documento"


def _blocks(md: str):
    """Linhas do markdown -> (tipo, texto). Tipos: h1 h2 h3 li ol p hr blank."""
    for line in md.replace("\r\n", "\n").split("\n"):
        s = line.rstrip()
        if not s.strip():
            yield "blank", ""
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", s.strip())
        if m:
            yield f"h{len(m.group(1))}", m.group(2).replace("**", "")
            continue
        if re.match(r"^\s*(-{3,}|_{3,}|\*{3,})\s*$", s):
            yield "hr", ""
            continue
        m = re.match(r"^\s*[-*•]\s+(.*)$", s)
        if m:
            yield "li", m.group(1)
            continue
        m = re.match(r"^\s*(\d+[.)])\s+(.*)$", s)
        if m:
            yield "ol", f"{m.group(1)} {m.group(2)}"
            continue
        yield "p", s.strip()


def _inline(text: str) -> str:
    t = re.sub(r"\[([^\]\n]+)\]\((https?://[^\s)]+)\)", r"\1 (\2)", text)
    return t.replace("`", "")


def _runs(text: str):
    """Trechos (texto, negrito)."""
    for part in re.split(r"(\*\*.+?\*\*)", _inline(text)):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            yield part[2:-2], True
        elif part:
            yield part, False


def to_docx(md: str, path: Path):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for kind, text in _blocks(md):
        if kind == "blank":
            continue
        if kind in ("h1", "h2", "h3"):
            doc.add_heading(_inline(text), level=int(kind[1]))
            continue
        if kind == "hr":
            doc.add_paragraph()
            continue
        p = doc.add_paragraph(style="List Bullet") if kind == "li" else doc.add_paragraph()
        for t, bold in _runs(text):
            p.add_run(t).bold = bold
    doc.save(path)


_LATIN = {"—": "-", "–": "-", "“": '"', "”": '"', "‘": "'", "’": "'", "•": "·", "…": "...", " ": " ",
          "→": "->", "←": "<-", "✓": "v", "✔": "v", "€": "EUR"}


def _latin(s: str) -> str:
    """As fontes padrão do PDF só têm Latin-1 (acentos do português incluídos); troca o resto."""
    s = "".join(_LATIN.get(c, c) for c in s)
    return s.encode("latin-1", "ignore").decode("latin-1")


def to_pdf(md: str, path: Path):
    from fpdf import FPDF
    pdf = FPDF(format="A4")
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(True, margin=16)
    pdf.add_page()
    w = pdf.epw

    def md_line(t: str) -> str:              # só **negrito**; -- e __ viram texto comum
        t = _latin(_inline(t))
        return re.sub(r"-{2,}", "-", t).replace("__", "_")

    for kind, text in _blocks(md):
        if kind == "blank":
            pdf.ln(2)
            continue
        if kind in ("h1", "h2", "h3"):
            size = {"h1": 17, "h2": 13.5, "h3": 12}[kind]
            pdf.ln(2 if kind != "h1" else 0)
            pdf.set_font("Helvetica", "B", size)
            pdf.multi_cell(w, size * 0.5, _latin(_inline(text)), new_x="LMARGIN", new_y="NEXT")
            if kind == "h2":
                y = pdf.get_y() + 0.8
                pdf.set_draw_color(160, 160, 160)
                pdf.line(pdf.l_margin, y, pdf.l_margin + w, y)
                pdf.ln(2)
            pdf.ln(1)
            continue
        if kind == "hr":
            y = pdf.get_y() + 1.5
            pdf.set_draw_color(160, 160, 160)
            pdf.line(pdf.l_margin, y, pdf.l_margin + w, y)
            pdf.ln(4)
            continue
        pdf.set_font("Helvetica", "", 10.5)
        if kind in ("li", "ol"):
            pdf.set_x(pdf.l_margin + 3)
            bullet = "· " if kind == "li" else ""
            pdf.multi_cell(w - 3, 5.2, bullet + md_line(text), markdown=True, new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.multi_cell(w, 5.2, md_line(text), markdown=True, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


@tool("documento_criar",
      "Cria um documento .docx (editável, padrão) ou .pdf a partir de markdown simples (# títulos, - listas, **negrito**). "
      "Use para CV, carta de apresentação, planos e relatórios. Devolve o caminho do arquivo; para mandar ao Bruno, "
      "use telegram_enviar com esse arquivo.",
      schema({"nome": ("string", "nome do arquivo sem extensão, ex.: CV Bruno - Empresa X"),
              "conteudo": ("string", "texto completo em markdown simples")},
             {"formato": ("string", "docx (padrão) | pdf | md")}))
@safe
def documento_criar(a):
    fmt = (a.get("formato") or "docx").lower().strip(". ")
    if fmt not in ("docx", "pdf", "md"):
        raise ValueError("formato deve ser docx, pdf ou md")
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / f"{_name(a['nome'])}.{fmt}"
    if fmt == "docx":
        to_docx(a["conteudo"], path)
    elif fmt == "pdf":
        to_pdf(a["conteudo"], path)
    else:
        path.write_text(a["conteudo"], encoding="utf-8")
    return f"Documento criado: {path}"


TOOLS = [documento_criar]
docs_server = create_sdk_mcp_server(name="docs", version="1.0.0", tools=TOOLS)
NAMES = [f"mcp__docs__{t.name}" for t in TOOLS]
